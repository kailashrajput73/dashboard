import unittest

from app.services.catalog_family_util import (
    build_family_key,
    canonicalize_pipes_master_taxonomy,
)
from app.services.pipes_fittings_import import parse_rows


class PipesCanonicalTaxonomyTests(unittest.TestCase):
    def test_category_type_class_brand_variants_one_family_key(self) -> None:
        variants = [
            ("pipes & fitting", "Cpvc", "Fitting ", " SDR 11", " king ", "COUPLER"),
            ("pipes & FITTING", "cPvc", "fITTING", "SDR   11", "KING", "coupler"),
            ("  Pipes & Fitting  ", " CPVC", "FITTING", "sdR 11", "  KING  ", "COUPLER"),
            ("PIPES & FITTING  ", "CPVC ", "fITTING", "SDR  11", "KING", "Coupler"),
        ]
        keys = []
        for args in variants:
            canon = canonicalize_pipes_master_taxonomy(
                category=args[0],
                type_name=args[1],
                subcategory=args[2],
                product_class=args[3],
                brand=args[4],
                product_family=args[5],
            )
            keys.append(
                build_family_key(
                    category=canon[0],
                    type_name=canon[1],
                    subcategory=canon[2],
                    product_class=canon[3],
                    brand=canon[4],
                    product_family=canon[5],
                )
            )
        self.assertEqual(len(set(keys)), 1)

    def test_canonical_display_values(self) -> None:
        cat, typ, sub, cls, brand, fam = canonicalize_pipes_master_taxonomy(
            category="pipes & fitting",
            type_name="cpvc",
            subcategory="fitting",
            product_class="sdr 11",
            brand="king",
            product_family="coupler",
        )
        self.assertEqual(cat, "Pipes & Fittings")
        self.assertEqual(typ, "CPVC")
        self.assertEqual(sub, "Fitting")
        self.assertEqual(cls, "SDR 11")
        self.assertEqual(brand, "King")
        self.assertEqual(fam, "COUPLER")

    def test_king_cpvc_fixture_sheet_single_family_in_file(self) -> None:
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        path = (
            "/home/ubuntu/.cursor/projects/workspace/uploads/"
            "bug_fix_check_king_cpvc_fitting_sheet_uploded_in_local_3__4e31.xlsx"
        )
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sheet1"]
        raw_rows: list[dict[str, object]] = []
        headers = [ws.cell(1, c).value for c in range(1, 21)]
        for r in range(2, ws.max_row + 1):
            row_dict: dict[str, object] = {}
            for c, header in enumerate(headers, start=1):
                if header:
                    row_dict[str(header)] = ws.cell(r, c).value
            raw_rows.append(row_dict)
        result = parse_rows(raw_rows)
        self.assertTrue(result.ok, result.errors[:3])
        self.assertEqual(len(result.rows), 193)
        self.assertEqual({row.product_class for row in result.rows}, {"SDR 11"})
        self.assertEqual({row.type_name for row in result.rows}, {"CPVC"})
        self.assertEqual({row.brand for row in result.rows}, {"King"})
        self.assertEqual({row.category for row in result.rows}, {"Pipes & Fittings"})
        # Many product_family names (COUPLER, REDUCER, …) → many families; not one blob.
        self.assertGreater(len({row.family_key for row in result.rows}), 1)


if __name__ == "__main__":
    unittest.main()
